# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
runlog.py — shared run-log helper for all Master Wave update scripts.

All scripts write to a single workbook: data/runlogs/master_runlog.xlsx (OneDrive)
Each script gets its own tab. Tabs are created automatically on first use.
New scripts just call runlog.append() with their tab name — no other setup needed.

Schema evolution: if a script adds new headers, they are appended as new columns
to the right of the existing ones. Values are mapped by header name (not position),
so old rows are unaffected and new rows always land in the correct column.

See also activity_log.py — a separate, simpler one-row-per-script-run log
("did it succeed, one-line summary") rather than this file's detailed
per-metric numeric tabs.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

import onedrive_paths

RUNLOG_PATH = onedrive_paths.RUNLOGS_DIR / "master_runlog.xlsx"


def append(tab_name: str, headers: list, row: list, path: Path = None) -> None:
    """
    Append one row to tab_name in a workbook (master_runlog.xlsx by default —
    pass `path` to reuse this same schema-evolving append logic against a
    different file, e.g. build_metrics_summary.py's metrics_summary.xlsx).

    Creates the workbook and/or tab (with bold headers) if they don't exist.
    If the tab already exists and new headers are present, they are added as
    new columns to the right — old data is preserved and unaffected.
    Values are always written by header name, never by position.

    Args:
        tab_name : display name for the sheet tab (e.g. "MVP Update")
        headers  : list of column header strings for this run
        row      : list of values matching headers (same order/length)
        path     : workbook path; defaults to RUNLOG_PATH
    """
    path = path or RUNLOG_PATH
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = openpyxl.load_workbook(path)
    else:
        wb = openpyxl.Workbook()
        for name in wb.sheetnames:
            del wb[name]

    if tab_name in wb.sheetnames:
        ws = wb[tab_name]

        # Read existing headers from row 1
        existing_headers = [cell.value for cell in ws[1]]

        # Add any new headers as additional columns to the right
        for h in headers:
            if h not in existing_headers:
                new_col = len(existing_headers) + 1
                cell = ws.cell(row=1, column=new_col, value=h)
                cell.font = Font(bold=True)
                existing_headers.append(h)

        # Map values to column positions by header name
        row_dict = dict(zip(headers, row))
        ordered_row = [row_dict.get(h) for h in existing_headers]
        ws.append(ordered_row)

    else:
        ws = wb.create_sheet(tab_name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.append(row)

    wb.save(path)
