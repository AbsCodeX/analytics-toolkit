# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
csi_tracker.py — maintain the rolling CSI interactions tracker.

Usage (from C:\\path\\to\\analytics):
    python scripts\\csi_tracker.py            # ingest every PDF/Excel in the inbox
    python scripts\\csi_tracker.py --dry-run  # parse + report, change nothing

Process: drop each new "CSI interactions by day" export (PDF or .xlsx) into
data\\raw\\csi_interactions\\, then run this script. It:
  1. parses every PDF (pdfplumber) and Excel file in the inbox. Excel comes in
     two shapes: a proper 9-column table, or the flattened one-column export
     (whole row as one string, descriptions/names truncated at line breaks,
     AM/PM missing). For the flattened shape the parser anchors on
     State/Contact-type values, infers AM/PM from ticket ordering, and
     reconstructs truncated caller names against HR — unresolvable names are
     flagged for review, never guessed silently,
  2. upserts rows into the tracker by IMS Number — new tickets are appended
     and enriched (hierarchy from report.users/report.hr, provisioned roles +
     title/BU/wave from MVP raw.mvp, keyword-based Issue Category); existing
     tickets only get State / Assigned to / Updated / Updated by refreshed,
  3. NEVER overwrites manually edited Issue Category or Notes cells,
  4. rebuilds Master File Only, Analysis, and Analysis - Master File sheets,
  5. logs the ingest on the Loads sheet and moves the PDF to archive\\<YYYY-MM>\\.

Close the tracker in Excel before running — the save will fail otherwise.
"""

import argparse
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

# re-exec under the analytics venv (pdfplumber lives there)
_VENV_PY = Path(__file__).resolve().parents[1] / "analytics_env" / "Scripts" / "python.exe"
if _VENV_PY.exists() and Path(sys.executable).resolve() != _VENV_PY.resolve():
    sys.exit(subprocess.call([str(_VENV_PY), *map(str, sys.argv)]))

import pdfplumber  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from onedrive_paths import CSI_ARCHIVE_DIR, CSI_INBOX_DIR, CSI_TRACKER_PATH, month_subdir  # noqa: E402

SQLCMD = ["sqlcmd", "-S", r".\SQLEXPRESS", "-d", "AnalyticsDB", "-E", "-C", "-W", "-s", "|"]

# Interactions sheet layout (tracker owns these positions)
COLS = ["Number", "Opened", "Short description", "Opened for", "State",
        "Contact type", "Assigned to", "Updated", "Updated by",
        "Universal ID", "Leader", "AVP", "VP", "SVP",
        "Job Role 1", "Job Role 2", "Job Role 3", "Job Role 4", "Match Source",
        "Issue Category", "Department", "MVP Job Title", "MVP Business Unit",
        "Go Live Wave", "Notes", "Trend Theme"]
WIDTHS = [13, 22, 60, 22, 17, 13, 20, 22, 13,
          14, 22, 22, 22, 18, 30, 30, 30, 30, 34, 26, 28, 34, 26, 12, 40, 32]
C = {name: i + 1 for i, name in enumerate(COLS)}
PDF_FIELDS = ["Number", "Opened", "Short description", "Opened for", "State",
              "Contact type", "Assigned to", "Updated", "Updated by"]
REFRESH_ON_UPDATE = ["State", "Assigned to", "Updated by"]

STATES = ["Closed Complete", "Closed Abandoned", "Closed Incomplete",
          "Work in Progress", "Resolved", "Open", "New"]
CONTACT_TYPES = ["Phone", "Chat", "Email", "Walk-in", "Self-service", "Video"]
FLAT_ROW_RE = re.compile(
    r"^(IMS\d+) (\d{2}/\d{2}/\d{4}) (\d{1,2}:\d{2}:\d{2}) (.+?) "
    rf"({'|'.join(STATES)}) ({'|'.join(CONTACT_TYPES)}) "
    r"(.+?) (\d{2}/\d{2}/\d{4}) (\d{1,2}:\d{2}:\d{2}) (\S+)$")

# first matching rule wins; checked against the lower-cased description
CATEGORY_RULES = [
    (r"^wave \d", None),  # placeholder — wave prefix handled explicitly below
    (r"call drop|dropped call|dead air|no answer|couldn.?t hear|could not hear|"
     r"unable to hear|^test\b", "Dropped / test call (no issue)"),
    (r"in ?basket|query|results? (aren't|not|notes)|recipient", "In Basket / messaging"),
    (r"follow.?up|following up|inc\d|ritm\d", "Ticket follow-up"),
    (r"telehealth|video visit|virtual visit", "Telehealth / video visits"),
    (r"e-?prescribe|placing order|order error|cosign required", "Orders / e-prescribe"),
    (r"template|schedule|appointment|booking|induction", "Scheduling / templates"),
    (r"sign (a )?(note|visit)|note|letter|telephone encounter|scanned|abridge|encounter opened",
     "Notes / documentation"),
    (r"lock(ed)? out|login|log in|unable to (login|access epic)|can ?not access epic|imprivata|"
     r"account (un)?blocked|hyperspace|session|password", "Login / account lockout"),
    (r"break the glass|needs access|have access|can.?t find workqueue|access to a workqueue|"
     r"careport|training env", "Access / security"),
    (r"registration|insurance|billing|charge|cash|refund|estimate|drg|coverage|copay",
     "Registration / billing"),
    (r"print|ekg|rover|locker|wrist ?band|device", "Devices / printing / hardware"),
    (r"haiku|canto", "Mobile (Haiku)"),
    (r"mychart|how to|wants to know|transcript|filter", "How-to / personalization"),
]


def categorize(desc):
    d = (desc or "").lower()
    m = re.match(r"^wave (\d)", d)
    if m:
        return f"Wave {m.group(1)} go-live support"
    for pattern, cat in CATEGORY_RULES:
        if cat and re.search(pattern, d):
            return cat
    return "REVIEW — categorize"


# Rollup to the trend themes leadership uses in their email summaries.
# ALWAYS derived from Issue Category + description — edit the Issue Category
# cell (not this column) to change a ticket's theme.
THEME_BY_CATEGORY = {
    "Access / security": "Access & Security",
    "Login / account lockout": "Access & Security",
    "Registration / billing": "Billing Errors & Revenue Cycle",
    "Scheduling / templates": "Workflow Support & How-To",
    "Notes / documentation": "Workflow Support & How-To",
    "Telehealth / video visits": "Workflow Support & How-To",
    "How-to / personalization": "Workflow Support & How-To",
    "In Basket / messaging": "Workflow Support & How-To",
    "Orders / e-prescribe": "Workflow Support & How-To",
    "Mobile (Haiku)": "Workflow Support & How-To",
    "Ticket follow-up": "Other — ticket follow-up",
    "Devices / printing / hardware": "Other — devices/hardware",
    "Provider record update": "Other — provider records",
    "Dropped / test call (no issue)": "Other — dropped/test calls",
}


def trend_theme(cat, desc):
    cat, d = cat or "", (desc or "").lower()
    if re.search(r"work ?queue|\bwq\b", d):
        return "Work Queue (WQ) Management"
    if cat in ("How-to / personalization", "Access / security") and re.search(r"training|transcript", d):
        return "Training & Knowledge Gaps"
    if cat.startswith("Wave"):
        return "Go-live support (Wave)"
    return THEME_BY_CATEGORY.get(cat, "Other")


# ---------------------------------------------------------------- PDF parsing
def clean_cell(v):
    if v is None:
        return ""
    v = v.replace("-\n", "-").replace("\n", " ")
    return re.sub(r"\s+", " ", v).strip()


def parse_pdf(path):
    """Return (meta dict, list of 9-field row dicts)."""
    rows, meta = [], {}
    with pdfplumber.open(path) as pdf:
        text1 = pdf.pages[0].extract_text() or ""
        for key in ["Run Date and Time", "Run by", "Query Condition"]:
            m = re.search(rf"{key}:\s*(.+)", text1)
            if m:
                meta[key] = m.group(1).strip()
        m = re.search(r"(\d+)\s+Interactions", text1)
        meta["Total"] = int(m.group(1)) if m else None
        for page in pdf.pages:
            for table in page.extract_tables():
                for raw in table:
                    if not raw or not raw[0] or not str(raw[0]).startswith("IMS"):
                        continue
                    vals = [clean_cell(c) for c in raw[:9]]
                    if len(vals) == 9:
                        rows.append(dict(zip(PDF_FIELDS, vals)))
    if meta.get("Total") is not None and meta["Total"] != len(rows):
        raise SystemExit(f"{path.name}: PDF says {meta['Total']} interactions "
                         f"but parsed {len(rows)} — layout changed? Aborting, nothing written.")
    return meta, rows


def parse_dt(s):
    return datetime.strptime(s, "%m/%d/%Y %I:%M:%S %p")


def to_dt(v):
    return v if isinstance(v, datetime) else parse_dt(v)


# ------------------------------------------------------------- Excel parsing
def _pick_time(date_s, time_s, floor):
    """Choose AM or PM for an hour with no meridiem: the earliest option that
    is >= floor and not before 7:00 AM (support desk hours)."""
    d = datetime.strptime(date_s, "%m/%d/%Y")
    h, mi, s = map(int, time_s.split(":"))
    options = sorted(d.replace(hour=(h % 12) + off, minute=mi, second=s) for off in (0, 12))
    good = [o for o in options if o >= floor and o.hour >= 7]
    return good[0] if good else options[-1]


def parse_xlsx(path):
    """Parse an Excel export. Returns (meta, rows) like parse_pdf. Rows from
    the flattened one-column shape carry a 'Blob' (desc+name merged) that
    resolve_blobs() must split before upsert."""
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    grid = [r for r in ws.iter_rows(values_only=True)
            if r and any(v is not None for v in r)]
    meta = {"Run Date and Time": None, "Query Condition": f"(Excel export: {path.name})", "Total": None}
    header = [str(v).strip() if v else "" for v in grid[0]]
    rows = []
    if header[: len(PDF_FIELDS)] == PDF_FIELDS and len(grid) > 1 and grid[1][1] is not None:
        # proper 9-column table
        for raw in grid[1:]:
            vals = list(raw[:9])
            if not vals[0] or not str(vals[0]).startswith("IMS"):
                continue
            rec = dict(zip(PDF_FIELDS, [v if isinstance(v, datetime) else clean_cell(str(v)) for v in vals]))
            rows.append(rec)
        return meta, rows
    # flattened one-column shape
    last_by_day = {}
    for raw in grid[1:]:
        v = clean_cell(str(raw[0] or ""))
        if not v.startswith("IMS"):
            continue
        m = FLAT_ROW_RE.match(v)
        if not m:
            raise SystemExit(f"{path.name}: could not parse flattened row — layout changed? "
                             f"Aborting, nothing written.\n  {v[:160]}")
        num, od, ot, blob, state, contact, assigned, ud, ut, upby = m.groups()
        day = datetime.strptime(od, "%m/%d/%Y").date()
        opened = _pick_time(od, ot, last_by_day.get(day, datetime.min))
        last_by_day[day] = opened
        rows.append({"Number": num, "Opened": opened, "Blob": blob, "State": state,
                     "Contact type": contact, "Assigned to": assigned,
                     "Updated": _pick_time(ud, ut, opened), "Updated by": upby})
    return meta, rows


def _reconstruct_name(person_row):
    last, given = person_row["FullName"].split(",", 1)
    return f"{given.strip()} {last.strip()}"


def resolve_blobs(rows, users_idx, hr_idx):
    """Split each Blob into Short description + Opened for. Truncated names
    (ending '-') are reconstructed against HR when the match is unique."""
    flagged = []
    for rec in rows:
        blob = rec.pop("Blob", None)
        if blob is None:
            continue
        m = re.search(r"Unlisted (YourOrg|Non)-?$", blob, re.IGNORECASE)
        if m:
            rec["Opened for"] = ("Unlisted YourOrg-Employee" if m.group(1).lower() == "yourorg"
                                 else "Unlisted Non-YourOrg")
            rec["Short description"] = blob[: m.start()].strip()
            continue
        toks = blob.split()
        resolved = None
        for L in range(min(5, len(toks) - 1), 0, -1):
            cand = " ".join(toks[-L:])
            if cand.endswith("-"):  # name truncated at a line-wrap hyphen
                given = [norm(t) for t in toks[-L:-1]]
                sp = norm(cand.split()[-1])
                hits = {}
                for idx in (users_idx, hr_idx):
                    for surname, lst in idx.items():
                        if not surname.startswith(sp):
                            continue
                        for gv, row in lst:
                            if not given or (gv and gv[0] == given[0] and gv[: len(given)] == given):
                                hits[row.get("UniversalID") or id(row)] = row
                if len(hits) == 1:
                    resolved = (blob[: -len(cand)].strip(), _reconstruct_name(next(iter(hits.values()))))
                    break
            elif L >= 2 and (find_matches(cand, users_idx) or find_matches(cand, hr_idx)):
                resolved = (blob[: -len(cand)].strip(), cand)
                break
        if resolved:
            rec["Short description"], rec["Opened for"] = resolved
        else:  # can't tell where the description ends and the name starts
            rec["Short description"] = " ".join(toks[:-2])
            rec["Opened for"] = " ".join(toks[-2:])
            rec["_name_flag"] = True
            flagged.append(rec["Number"])
    return flagged


# ---------------------------------------------------------------- SQL lookups
def sql_rows(query):
    """Run a query via sqlcmd, return list of dicts (NULL -> None)."""
    out = subprocess.run(SQLCMD + ["-Q", "SET NOCOUNT ON; " + query],
                         capture_output=True, text=True, check=True).stdout
    lines = [ln for ln in out.splitlines() if ln.strip()]
    if len(lines) < 2:
        return []
    header = [h.strip() for h in lines[0].split("|")]
    rows = []
    for ln in lines[2:]:
        parts = [p.strip() for p in ln.split("|")]
        if len(parts) == len(header):
            rows.append({h: (None if v == "NULL" else v) for h, v in zip(header, parts)})
    return rows


def norm(s):
    s = (s or "").lower().replace("’", "")
    s = re.sub(r"[.']", "", s)
    return re.sub(r"\s+", " ", s.replace("-", " ")).strip()


def name_index(rows):
    idx = defaultdict(list)
    for r in rows:
        fn = r.get("FullName") or ""
        if "," not in fn:
            continue
        last, given = fn.split(",", 1)
        idx[norm(last)].append((norm(given).split(), r))
    return idx


def find_matches(pdf_name, idx):
    toks = norm(pdf_name).split()
    hits = {}
    for k in range(1, len(toks)):
        given, surname = toks[:k], " ".join(toks[k:])
        for gv, row in idx.get(surname, []):
            if gv and gv[0] == given[0] and (gv[: len(given)] == given or len(given) == 1):
                hits[row.get("UniversalID") or id(row)] = row
    return list(hits.values())


def load_reference_data():
    users = sql_rows("SELECT UniversalID, FullName, Leader, AVP, VP, SVP, "
                     "JobRole1, JobRole2, JobRole3, JobRole4, Departed FROM report.users")
    hr = sql_rows("SELECT UniversalID, FullName, Leader, AVP, VP, SVP, "
                  "Department, JobTitle, Departed FROM report.hr")
    return name_index(users), name_index(hr), {r["UniversalID"].upper(): r for r in hr if r.get("UniversalID")}


def mvp_lookup(uids):
    if not uids:
        return {}
    ids = ",".join(f"'{u}'" for u in sorted(uids))
    rows = sql_rows("SELECT UniversalID, IndividualCategoryUpdate1Name, IndividualCategoryUpdate2Name, "
                    "IndividualCategoryUpdate3Name, IndividualCategoryUpdate4Name, JobTitle, "
                    "BusinessUnit, BusinessUnitDescription, GoLiveWave "
                    f"FROM raw.mvp WHERE UniversalID IN ({ids})")
    return {r["UniversalID"].upper(): r for r in rows}


def resolve_person(name, users_idx, hr_idx, mvp_by_uid_probe):
    """Match one 'First Last' name. Returns (row, source_label, is_wave)."""
    if norm(name) == "unlisted yourorg employee":
        return None, "N/A — unlisted employee placeholder", False
    m = find_matches(name, users_idx)
    if m:
        active = [x for x in m if (x.get("Departed") or "").lower() != "true"] or m
        note = "" if len(active) == 1 else " (multiple wave matches; first used)"
        return active[0], "Master wave file" + note, True
    m = find_matches(name, hr_idx)
    if not m:
        return None, "Not found in wave file or HR", False
    active = [x for x in m if (x.get("Departed") or "").lower() != "true"] or m
    if len(active) > 1:
        # prefer the candidate with a real Epic provisioning row in MVP
        probe = mvp_by_uid_probe({x["UniversalID"].upper() for x in active if x.get("UniversalID")})
        real = [x for x in active
                if (probe.get((x.get("UniversalID") or "").upper(), {})
                    .get("IndividualCategoryUpdate1Name") or "").strip()
                not in ("", "MAPPING IN PROGRESS", "No Access or Training Needed")]
        if len(real) == 1:
            return real[0], "HR file — not on wave file (same-name matches; resolved via MVP provisioning)", False
        ids = ", ".join(sorted((x.get("UniversalID") or "?") for x in active))
        return active[0], f"HR file — not on wave file; AMBIGUOUS ({ids}) — verify", False
    return active[0], "HR file — not on wave file", False


# ---------------------------------------------------------------- workbook
def styled_header(ws, names, fill_hex):
    for i, h in enumerate(names, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = PatternFill("solid", fgColor=fill_hex)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center")


def ensure_layout(wb):
    """Migrate older trackers in place: add any new columns / sheets."""
    ws = wb["Interactions"]
    for i, name in enumerate(COLS, 1):
        if ws.cell(row=1, column=i).value != name:
            c = ws.cell(row=1, column=i, value=name)
            c.fill = PatternFill("solid", fgColor="5B9BD5")
            c.font = Font(bold=True, color="FFFFFF")
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
    if "Changes" not in wb.sheetnames:
        ch = wb.create_sheet("Changes")
        styled_header(ch, ["Ingested at", "File", "Number", "Action", "Opened for", "Details"], "ED7D31")
        for col, w in zip("ABCDEF", [20, 45, 13, 10, 22, 80]):
            ch.column_dimensions[col].width = w
    return wb


def open_tracker():
    if CSI_TRACKER_PATH.exists():
        return ensure_layout(load_workbook(CSI_TRACKER_PATH))
    CSI_TRACKER_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Interactions"
    styled_header(ws, COLS, "5B9BD5")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    loads = wb.create_sheet("Loads")
    styled_header(loads, ["Ingested at", "File", "Report run", "Query condition",
                          "Rows in PDF", "Added", "Updated"], "7030A0")
    for col, w in zip("ABCDEFG", [20, 52, 42, 90, 12, 10, 10]):
        loads.column_dimensions[col].width = w
    return ensure_layout(wb)


def upsert(wb, meta, pdf_rows, refs):
    ws = wb["Interactions"]
    users_idx, hr_idx, hr_by_uid = refs
    existing = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    changes = wb["Changes"]
    stamp = datetime.now().strftime("%m/%d/%Y %I:%M %p")
    added = updated = 0
    new_rows = []
    for rec in pdf_rows:
        num = rec["Number"]
        if num in existing:
            r = existing[num]
            diffs = []
            for f in REFRESH_ON_UPDATE:
                new_v, old_v = rec[f], ws.cell(row=r, column=C[f]).value
                if old_v != new_v:
                    ws.cell(row=r, column=C[f], value=new_v)
                    diffs.append(f"{f}: {old_v} → {new_v}")
            if diffs:
                # refresh the Updated timestamp only alongside a real change —
                # degraded sources (flattened Excel) carry inferred times
                cell = ws.cell(row=r, column=C["Updated"], value=to_dt(rec["Updated"]))
                cell.number_format = "mm/dd/yyyy hh:mm:ss AM/PM"
                changes.append([stamp, meta["file"], num, "Updated", rec["Opened for"], "; ".join(diffs)])
                updated += 1
            continue
        r = ws.max_row + 1
        existing[num] = r
        for f in PDF_FIELDS:
            v = to_dt(rec[f]) if f in ("Opened", "Updated") else rec[f]
            cell = ws.cell(row=r, column=C[f], value=v)
            if f in ("Opened", "Updated"):
                cell.number_format = "mm/dd/yyyy hh:mm:ss AM/PM"
            cell.alignment = Alignment(vertical="top", wrap_text=(f == "Short description"))
        person, source, _ = resolve_person(rec["Opened for"], users_idx, hr_idx, mvp_lookup)
        if rec.get("_name_flag"):
            source += " — NAME MAY BE TRUNCATED in Excel export, verify caller"
        uid = (person or {}).get("UniversalID")
        hr_row = hr_by_uid.get((uid or "").upper(), {})
        for f, v in [("Universal ID", uid), ("Leader", (person or {}).get("Leader")),
                     ("AVP", (person or {}).get("AVP")), ("VP", (person or {}).get("VP")),
                     ("SVP", (person or {}).get("SVP")), ("Match Source", source),
                     ("Department", hr_row.get("Department")),
                     ("Issue Category", categorize(rec["Short description"]))]:
            ws.cell(row=r, column=C[f], value=v).alignment = Alignment(vertical="top")
        for i in range(1, 5):  # Master job roles for wave people
            ws.cell(row=r, column=C[f"Job Role {i}"],
                    value=(person or {}).get(f"JobRole{i}")).alignment = Alignment(vertical="top")
        changes.append([stamp, meta["file"], num, "Added", rec["Opened for"],
                        f'{to_dt(rec["Opened"]):%m/%d/%Y} — {rec["Short description"][:120]}'])
        new_rows.append(r)
        added += 1
    # MVP fill for the new rows (roles where blank + title/BU/wave)
    uids = {ws.cell(row=r, column=C["Universal ID"]).value for r in new_rows}
    mvp = mvp_lookup({u.upper() for u in uids if u})
    for r in new_rows:
        uid = (ws.cell(row=r, column=C["Universal ID"]).value or "").upper()
        m = mvp.get(uid)
        if not m:
            continue
        if not ws.cell(row=r, column=C["Job Role 1"]).value:
            for i in range(1, 5):
                ws.cell(row=r, column=C[f"Job Role {i}"],
                        value=m.get(f"IndividualCategoryUpdate{i}Name")).alignment = Alignment(vertical="top")
            src = ws.cell(row=r, column=C["Match Source"])
            if src.value and "MVP" not in src.value:
                src.value += "; job roles from MVP user mappings"
        bu = m.get("BusinessUnit")
        if bu and m.get("BusinessUnitDescription"):
            bu = f"{bu} — {m['BusinessUnitDescription']}"
        for f, v in [("MVP Job Title", m.get("JobTitle")), ("MVP Business Unit", bu),
                     ("Go Live Wave", m.get("GoLiveWave"))]:
            ws.cell(row=r, column=C[f], value=v).alignment = Alignment(vertical="top")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"
    loads = wb["Loads"]
    loads.append([datetime.now().strftime("%m/%d/%Y %I:%M %p"), meta["file"],
                  meta.get("Run Date and Time"), meta.get("Query Condition"),
                  len(pdf_rows), added, updated])
    return added, updated


# ------------------------------------------------------- analysis sheets
def read_rows(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        def g(name):
            return ws.cell(row=r, column=C[name]).value
        if not g("Number"):
            continue
        rows.append({
            "num": g("Number"), "day": g("Opened").strftime("%m/%d"), "who": g("Opened for"),
            "desc": g("Short description"), "state": g("State"), "contact": g("Contact type"),
            "assigned": g("Assigned to"), "leader": g("Leader"), "src": g("Match Source") or "",
            "cat": g("Issue Category") or "(blank)", "theme": g("Trend Theme") or "Other",
            "dept": g("Department") or "(unknown)",
            "title": g("MVP Job Title") or "(unknown)", "bu": g("MVP Business Unit") or "(none)",
            "wave": g("Go Live Wave") or "(none)",
        })
    return rows


def build_analysis(wb, sheet_name, rows, scope_label, days):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    an = wb.create_sheet(sheet_name)
    for col, w in zip("ABCDEF", [44, 10, 36, 70, 18, 18]):
        an.column_dimensions[col].width = w
    r = 1

    def title(t):
        nonlocal r
        an.cell(row=r, column=1, value=t).font = Font(bold=True, size=12)
        r += 1

    def table(headers, data, wrap_col=None):
        nonlocal r
        for i, h in enumerate(headers):
            c = an.cell(row=r, column=1 + i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="5B9BD5")
        r += 1
        if not data:
            an.cell(row=r, column=1, value="(none)")
            r += 1
        for row in data:
            for i, v in enumerate(row):
                cell = an.cell(row=r, column=1 + i, value=v)
                cell.alignment = Alignment(vertical="top", wrap_text=(i + 1 == wrap_col))
            r += 1
        r += 1

    span = f"{days[0]}–{days[-1]}" if days else ""
    an.cell(row=r, column=1, value=f"CSI Interactions Analysis — {scope_label} — {span}").font = Font(bold=True, size=14)
    r += 1
    an.cell(row=r, column=1, value="TOTAL CALLS").font = Font(bold=True)
    an.cell(row=r, column=2, value=len(rows)).font = Font(bold=True)
    r += 2
    title("Call volume by day")
    table(["Day", "Calls"], sorted(Counter(x["day"] for x in rows).items()))
    title("Call volume by leader")
    table(["Leader", "Calls"], Counter(x["leader"] or "(none listed)" for x in rows).most_common())
    title("Trend themes by day (leadership rollup)")
    themeday = defaultdict(Counter)
    for x in rows:
        themeday[x["theme"]][x["day"]] += 1
    table(["Trend Theme", "Total"] + days,
          [(t, sum(c.values()), *(c[d] for d in days))
           for t, c in sorted(themeday.items(), key=lambda kv: -sum(kv[1].values()))])
    title("Issue categories by day")
    byday = defaultdict(Counter)
    for x in rows:
        byday[x["cat"]][x["day"]] += 1
    table(["Category", "Total"] + days,
          [(cat, sum(c.values()), *(c[d] for d in days))
           for cat, c in sorted(byday.items(), key=lambda kv: -sum(kv[1].values()))])
    rep = Counter(x["who"] for x in rows)
    title("Total calls by caller")
    table(["Caller", "Calls", "Job Title", "Department", "Categories"],
          [(w, n, next(x["title"] for x in rows if x["who"] == w),
            next(x["dept"] for x in rows if x["who"] == w),
            "; ".join(sorted({x["cat"] for x in rows if x["who"] == w})))
           for w, n in rep.most_common()])
    title("Repeat callers (2+ contacts)")
    table(["Caller", "Calls", "Title / Department", "What they called about"],
          [(w, n,
            f'{next(x["title"] for x in rows if x["who"] == w)} — {next(x["dept"] for x in rows if x["who"] == w)}',
            "\n".join(f'{x["num"]} ({x["day"]}) [{x["cat"]}] {x["desc"]}' for x in rows if x["who"] == w))
           for w, n in rep.most_common() if n > 1], wrap_col=4)
    title("Call volume by department")
    table(["Department", "Calls"], Counter(x["dept"] for x in rows).most_common())
    title("Call volume by go-live wave")
    table(["Go Live Wave", "Calls"], Counter(x["wave"] for x in rows).most_common())
    title("Call volume by business unit")
    table(["Business Unit", "Calls"], Counter(x["bu"] for x in rows).most_common())
    title("State")
    table(["State", "Calls"], Counter(x["state"] for x in rows).most_common())
    title("Contact type")
    table(["Contact type", "Calls"], Counter(x["contact"] for x in rows).most_common())
    title("Assigned to")
    table(["Assigned to", "Calls"], Counter(x["assigned"] for x in rows).most_common())
    an.cell(row=r, column=1, value="TOTAL CALLS").font = Font(bold=True)
    an.cell(row=r, column=2, value=len(rows)).font = Font(bold=True)


def build_master_sheet(wb, all_rows):
    if "Master File Only" in wb.sheetnames:
        del wb["Master File Only"]
    src = wb["Interactions"]
    ws = wb.create_sheet("Master File Only")
    styled_header(ws, COLS + ["Total Calls by Caller"], "C00000")
    for i, w in enumerate(WIDTHS + [18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    totals = Counter(x["who"] for x in all_rows)
    out = 2
    for r in range(2, src.max_row + 1):
        if not (src.cell(row=r, column=C["Match Source"]).value or "").startswith("Master wave file"):
            continue
        for c in range(1, len(COLS) + 1):
            s = src.cell(row=r, column=c)
            d = ws.cell(row=out, column=c, value=s.value)
            d.number_format = s.number_format
            d.alignment = Alignment(vertical="top", wrap_text=(c == C["Short description"]))
        ws.cell(row=out, column=len(COLS) + 1,
                value=totals[src.cell(row=r, column=C["Opened for"]).value])
        out += 1
    ws.cell(row=out, column=1, value="TOTAL CALLS").font = Font(bold=True)
    ws.cell(row=out, column=2, value=out - 2).font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS) + 1)}{out - 1}"


def rebuild_outputs(wb):
    ws = wb["Interactions"]
    for r in range(2, ws.max_row + 1):  # Trend Theme is always derived
        if ws.cell(row=r, column=1).value:
            ws.cell(row=r, column=C["Trend Theme"],
                    value=trend_theme(ws.cell(row=r, column=C["Issue Category"]).value,
                                      ws.cell(row=r, column=C["Short description"]).value)
                    ).alignment = Alignment(vertical="top")
    all_rows = read_rows(wb["Interactions"])
    days = sorted({x["day"] for x in all_rows})
    master_rows = [x for x in all_rows if x["src"].startswith("Master wave file")]
    build_master_sheet(wb, all_rows)
    build_analysis(wb, "Analysis", all_rows, f"All Interactions ({len(all_rows)})", days)
    build_analysis(wb, "Analysis - Master File", master_rows,
                   f"Master Wave File Population Only ({len(master_rows)} of {len(all_rows)})", days)
    order = ["Interactions", "Master File Only", "Analysis", "Analysis - Master File",
             "Changes", "Loads"]
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
    return all_rows, master_rows


def day_over_day(all_rows):
    """Console comparison of the two most recent report days."""
    days = sorted({x["day"] for x in all_rows})
    if len(days) < 2:
        return
    prev, last = days[-2], days[-1]
    a = [x for x in all_rows if x["day"] == prev]
    b = [x for x in all_rows if x["day"] == last]
    print(f"\n=== Day over day: {prev} ({len(a)} calls) vs {last} ({len(b)} calls) ===")
    ta, tb = Counter(x["theme"] for x in a), Counter(x["theme"] for x in b)
    for theme in sorted(set(ta) | set(tb), key=lambda t: -(tb[t])):
        delta = tb[theme] - ta[theme]
        print(f"  {theme:38} {prev}: {ta[theme]:2}   {last}: {tb[theme]:2}   ({delta:+d})")
    prev_names = {x["who"] for x in a}
    both = sorted({x["who"] for x in b} & prev_names)
    if both:
        print(f"  Called on both days: {', '.join(both)}")


def main():
    ap = argparse.ArgumentParser(description="Merge CSI interaction PDFs into the rolling tracker")
    ap.add_argument("--dry-run", action="store_true", help="parse and report only; write/move nothing")
    ap.add_argument("--rebuild", action="store_true",
                    help="rebuild analysis sheets from the tracker even with no new PDFs")
    args = ap.parse_args()

    CSI_INBOX_DIR.mkdir(parents=True, exist_ok=True)
    inbox = sorted(p for p in CSI_INBOX_DIR.iterdir()
                   if p.suffix.lower() in (".pdf", ".xlsx") and not p.name.startswith("~$"))
    if not inbox and not args.rebuild:
        print(f"No PDF/Excel files in {CSI_INBOX_DIR} — nothing to do. "
              f"(--rebuild to refresh sheets anyway)")
        return
    wb = open_tracker()
    refs = load_reference_data()
    pre_nums = {wb["Interactions"].cell(row=r, column=1).value
                for r in range(2, wb["Interactions"].max_row + 1)}
    name_flags = []
    for src in inbox:
        meta, rows = parse_pdf(src) if src.suffix.lower() == ".pdf" else parse_xlsx(src)
        meta["file"] = src.name
        # flags only matter on rows actually entering the tracker
        name_flags += [n for n in resolve_blobs(rows, refs[0], refs[1]) if n not in pre_nums]
        if args.dry_run:
            nums = {wb["Interactions"].cell(row=r, column=1).value
                    for r in range(2, wb["Interactions"].max_row + 1)}
            new = [x for x in rows if x["Number"] not in nums]
            print(f"[dry-run] {src.name}: {len(rows)} rows, {len(new)} new")
            continue
        added, updated = upsert(wb, meta, rows, refs)
        print(f"{src.name}: {len(rows)} rows parsed — {added} added, {updated} updated")
    if name_flags:
        print(f"Caller names needing verification (truncated in Excel export): {', '.join(name_flags)}")
    if args.dry_run:
        return
    all_rows, master_rows = rebuild_outputs(wb)
    wb.save(CSI_TRACKER_PATH)
    for src in inbox:
        shutil.move(str(src), month_subdir(CSI_ARCHIVE_DIR) / src.name)
    review = [x["num"] for x in all_rows if x["cat"].startswith("REVIEW")]
    ambiguous = [x["num"] for x in all_rows if "AMBIGUOUS" in x["src"] or "TRUNCATED" in x["src"]]
    day_over_day(all_rows)
    print(f"\nTracker: {CSI_TRACKER_PATH}")
    print(f"Total {len(all_rows)} interactions | master file {len(master_rows)} | "
          f"repeat callers {sum(1 for _, n in Counter(x['who'] for x in all_rows).items() if n > 1)}")
    if review:
        print(f"Needs category review ({len(review)}): {', '.join(review)}")
    if ambiguous:
        print(f"Ambiguous name matches to verify: {', '.join(ambiguous)}")


if __name__ == "__main__":
    main()
