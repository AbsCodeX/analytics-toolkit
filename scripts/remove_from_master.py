# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
remove_from_master.py — remove specific people from the Master Wave File.

Her ask 2026-09-15 ("remove these people of the wave file"): a sanctioned,
audited way to delete rows from the DATA table for an explicit list of
Universal IDs. Nothing else in the pipeline deletes Master rows.

SAFE-WRITE PATTERN (same as add_missing_to_master.py):
  * Master is READ with openpyxl (from a temp copy if Excel holds the file) but
    only ever WRITTEN through desktop Excel (xlwings/COM) via the DATA sheet's
    ListObject ("Table1") ListRows(i).Delete(), so the table stays intact.
  * Daily full backup + values-only DATA snapshot via master_backup.
  * Refuses to write if the Master is open/locked in Excel.
  * DRY RUN by default. A review CSV holding the FULL row of every person
    removed is written in both modes (Main Reports/removed_from_master/), so a
    removal can be undone by re-adding the row.
  * Each run is logged to the Master runlog workbook ("Removed From Master" tab).

USAGE
-----
    python scripts/remove_from_master.py --uids USER06UID USER03UID --reason "not rev cycle (her call)"
    python scripts/remove_from_master.py --uids USER06UID USER03UID --reason "..." --apply
"""
import argparse
import datetime
import shutil
import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl  # READS only — never used to write the Master
import pandas as pd

warnings.filterwarnings("ignore")
sys.path.insert(0, str(Path(__file__).parent))
import master_backup
import onedrive_paths as op
import runlog
from activity_log import track_run

TODAY = datetime.date.today()
MASTER_PATH = op.MASTER_WAVE_PATH
DATA_SHEET = "DATA"
TABLE_NAME = "Table1"
CSV_DIR = MASTER_PATH.parent / "removed_from_master"
RUNLOG_TAB = "Removed From Master"
RUNLOG_HEADERS = ["Run Timestamp", "Reason", "UIDs Requested", "Not Found",
                  "Rows Removed", "UIDs Removed", "Review CSV", "Master File", "Backup File"]
SANITY_MAX_REMOVALS = 50


def office_lock_present(path: Path) -> bool:
    return path.with_name("~$" + path.name).exists()


def read_master_rows(uids: set) -> tuple[pd.DataFrame, list]:
    """Full DATA rows for the requested UIDs (values only), read from a temp
    copy so an open Excel session cannot block the preview."""
    tmp = Path(tempfile.gettempdir()) / f"master_read_{TODAY.isoformat()}.xlsx"
    shutil.copy2(MASTER_PATH, tmp)
    wb = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
    ws = wb[DATA_SHEET]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    uid_ix = headers.index("UniversalID")
    found, sheet_rows = [], []
    for r_no, row in enumerate(rows, start=2):
        u = str(row[uid_ix] or "").strip().upper()
        if u in uids:
            found.append(dict(zip(headers, row)))
            sheet_rows.append(r_no)
    wb.close()
    tmp.unlink(missing_ok=True)
    df = pd.DataFrame(found)
    if len(df):
        df.insert(0, "_sheet_row", sheet_rows)
    return df, headers


def main(uids: list, reason: str, apply: bool) -> None:
    banner = "LIVE RUN — WILL DELETE ROWS" if apply else "DRY RUN — NOTHING WILL BE WRITTEN"
    print("=" * 72)
    print(f"  remove_from_master.py   [{banner}]")
    print("=" * 72 + "\n")
    wanted = {str(u).strip().upper() for u in uids if str(u).strip()}
    df, headers = read_master_rows(wanted)
    found = set(df["UniversalID"].astype(str).str.upper()) if len(df) else set()
    missing = sorted(wanted - found)
    print(f"Requested: {len(wanted)} | found on Master: {len(found)} | not found: {len(missing)}")
    if missing:
        print(f"  not on the Master (skipped): {', '.join(missing)}")
    if not len(df):
        print("Nothing to remove.")
        return
    if len(df) > SANITY_MAX_REMOVALS:
        sys.exit(f"SANITY ABORT: {len(df)} rows exceeds {SANITY_MAX_REMOVALS}. No changes made.")

    show = [c for c in ["UniversalID", "Full Name", "Leaders", "GoLiveWave", "JobTitle",
                        "Training Needed (Yes/No)", "Notes"] if c in df.columns]
    pd.set_option("display.max_rows", None, "display.width", 220, "display.max_colwidth", 40)
    print("\nRows to remove:")
    print(df[show].to_string(index=False))

    CSV_DIR.mkdir(parents=True, exist_ok=True)
    csv_out = CSV_DIR / f"Removed_From_Master_{TODAY.isoformat()}_{datetime.datetime.now():%H%M%S}.csv"
    out = df.drop(columns=["_sheet_row"]).copy()
    out.insert(0, "Removal Reason", reason)
    out.insert(0, "Removed On", TODAY.isoformat())
    out.to_csv(csv_out, index=False)
    print(f"\n  Review CSV (full rows, for undo): {csv_out}")

    if not apply:
        print("\nDRY RUN — no backup made, nothing written. Re-run with --apply.")
        return

    import xlwings as xw

    if office_lock_present(MASTER_PATH):
        sys.exit("ERROR: Master Wave File is open/locked in Excel. Close it and re-run.")

    backup_path = master_backup.ensure_daily_full_backup(MASTER_PATH)
    print(f"Backup: {backup_path}")

    app = wb = None
    orig_calc = orig_alerts = orig_screen = None
    removed = []
    try:
        app = xw.App(visible=False, add_book=False)
        orig_alerts, orig_screen = app.display_alerts, app.screen_updating
        app.display_alerts = app.screen_updating = False
        wb = app.books.open(str(MASTER_PATH), update_links=False)
        orig_calc = app.calculation
        app.calculation = "manual"
        ws = wb.sheets[DATA_SHEET]
        lo = ws.tables[TABLE_NAME].api
        header_cells = lo.HeaderRowRange.Value[0]
        col_idx = {str(h).strip(): i for i, h in enumerate(header_cells) if h is not None}
        uid_col = col_idx["UniversalID"]
        # Excel cannot delete rows from a filtered table ("Can't move cells in a
        # filtered range or table"), so clear any active table filter first.
        # A filter is view state, not data; the saved file simply shows all rows.
        try:
            if lo.ShowAutoFilter and lo.AutoFilter.FilterMode:
                lo.AutoFilter.ShowAllData()
                print("Cleared the DATA table's active filter before deleting.")
        except Exception:
            pass
        body = lo.DataBodyRange.Value          # tuple of row tuples, live from Excel
        targets = []                           # (table row index 1-based, uid)
        for i, row in enumerate(body, start=1):
            u = str(row[uid_col] or "").strip().upper()
            if u in found:
                targets.append((i, u))
        if len(targets) != len(found):
            sys.exit(f"ERROR: expected {len(found)} table row(s), located {len(targets)} — "
                     "Master changed since the preview. Nothing deleted.")
        # delete bottom-up so earlier indexes stay valid
        for i, u in sorted(targets, reverse=True):
            lo.ListRows(i).Delete()
            removed.append(u)
        app.calculation = orig_calc
        orig_calc = None
        wb.save()
        print(f"Deleted {len(removed)} row(s); saved through Excel (workbook package + Table preserved).")
    finally:
        try:
            if app is not None and orig_calc is not None:
                app.calculation = orig_calc
        except Exception:
            pass
        try:
            if wb is not None:
                wb.close()
        except Exception:
            pass
        try:
            if app is not None:
                app.display_alerts = True if orig_alerts is None else orig_alerts
                app.screen_updating = True if orig_screen is None else orig_screen
                app.quit()
        except Exception:
            pass

    snap = master_backup.save_data_snapshot(MASTER_PATH)
    print(f"DATA snapshot: {snap}")
    runlog.append(RUNLOG_TAB, RUNLOG_HEADERS, [
        datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), reason,
        ", ".join(sorted(wanted)), ", ".join(missing), len(removed),
        ", ".join(sorted(removed)), str(csv_out), str(MASTER_PATH), str(backup_path),
    ])
    print(f"\nRemoved {len(removed)} member(s): {', '.join(sorted(removed))}")
    print("Next: apply --only sql_post (and lava / lava_workbook, tracker) so reports drop them.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--uids", nargs="+", required=True, metavar="UID")
    ap.add_argument("--reason", required=True, help="Short reason, recorded on the review CSV and runlog.")
    ap.add_argument("--apply", action="store_true", help="Back up the Master and delete. Default is a dry-run preview.")
    args = ap.parse_args()
    with track_run("Remove From Master"):
        main(args.uids, args.reason, args.apply)
