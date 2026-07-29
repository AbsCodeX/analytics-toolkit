# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
activity_log.py

High-level, one-row-per-script-run activity log — separate from runlog.py
(which stays as the detailed numeric per-metric log for the Master-update
scripts). This one is deliberately simple: every migrated script gets one row
telling you it ran, whether it succeeded, and a one-line summary, so there's a
single place to see "what did I run today and did it work."

Usage — wrap a script's entry point:

    if __name__ == "__main__":
        from activity_log import track_run
        with track_run("clean_hr.py"):
            main()

track_run tees stdout (prints normally to the console AND captures it), then
on success logs the last non-blank printed line as the summary — every
migrated script already ends with a good closing summary line, so no
per-script rewrite is needed to get a useful one-line description.
"""

from __future__ import annotations

import sys
from contextlib import contextmanager
from datetime import datetime

import openpyxl
from openpyxl.styles import Font

import onedrive_paths

LOG_PATH = onedrive_paths.RUNLOGS_DIR / "activity_log.xlsx"
SHEET_NAME = "Activity"
HEADERS = ["Timestamp", "Script", "Status", "Summary"]


class _Tee:
    """Writes to the real stream and a buffer at once, so wrapping a script
    with track_run doesn't change what the user sees in real time."""

    def __init__(self, real):
        self._real = real
        self.lines: list[str] = []
        self._partial = ""

    def write(self, s: str):
        self._real.write(s)
        self._partial += s
        while "\n" in self._partial:
            line, self._partial = self._partial.split("\n", 1)
            self.lines.append(line)
        return len(s)

    def flush(self):
        self._real.flush()


def _last_summary_line(lines: list[str]) -> str:
    for line in reversed(lines):
        if line.strip():
            return line.strip()[:300]
    return ""


def log_run(script_name: str, status: str, summary: str) -> None:
    """Append one row to activity_log.xlsx. Creates the file/sheet on first use.

    Best-effort by design: this wraps every script's entry point, so a logging
    failure (e.g. activity_log.xlsx open in Excel, OneDrive lock) must never
    change the wrapped script's real outcome. On any error it warns to the real
    stderr and returns rather than raising."""
    try:
        LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
        if LOG_PATH.exists():
            wb = openpyxl.load_workbook(LOG_PATH)
        else:
            wb = openpyxl.Workbook()
            for name in wb.sheetnames:
                del wb[name]

        if SHEET_NAME in wb.sheetnames:
            ws = wb[SHEET_NAME]
        else:
            ws = wb.create_sheet(SHEET_NAME)
            ws.append(HEADERS)
            for cell in ws[1]:
                cell.font = Font(bold=True)

        ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), script_name, status, summary])
        wb.save(LOG_PATH)
    except Exception as e:
        print(f"WARNING: could not write to activity_log ({type(e).__name__}: {e}).",
              file=sys.__stderr__)


@contextmanager
def track_run(script_name: str):
    tee_out = _Tee(sys.stdout)
    tee_err = _Tee(sys.stderr)
    orig_out, orig_err = sys.stdout, sys.stderr
    sys.stdout, sys.stderr = tee_out, tee_err
    try:
        yield
    except SystemExit as e:
        sys.stdout, sys.stderr = orig_out, orig_err
        code = e.code
        if code in (None, 0):
            log_run(script_name, "Success", _last_summary_line(tee_out.lines))
        elif isinstance(code, str) and code.strip():
            log_run(script_name, "Needs Review", code.strip()[:300])
        else:
            # Bare exit code (e.g. raise SystemExit(1)) isn't informative on its
            # own — the real message was almost certainly printed just before.
            msg = (_last_summary_line(tee_err.lines)
                   or _last_summary_line(tee_out.lines) or str(code))
            log_run(script_name, "Needs Review", msg)
        raise
    except Exception as e:
        sys.stdout, sys.stderr = orig_out, orig_err
        log_run(script_name, "Failed", f"{type(e).__name__}: {e}"[:300])
        raise
    else:
        sys.stdout, sys.stderr = orig_out, orig_err
        log_run(script_name, "Success", _last_summary_line(tee_out.lines))
