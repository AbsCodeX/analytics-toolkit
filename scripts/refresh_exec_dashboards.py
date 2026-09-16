# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
refresh_exec_dashboards.py

Stamp the RCM Executive Dashboard HTML files from the supplemental workbook.

Source of truth: Main Reports\\Dashboards\\RCM Dashboard Data.xlsx
(sheet "DashboardData" — one column of values per dashboard). Edit the
workbook, save it, then run:

    python scripts\\refresh_exec_dashboards.py

What it does, per dashboard file:
  - writes every metric into the HTML input defaults (the numbers the donuts
    render from on open)
  - picks the Training Completed denominator: the custom value if given,
    otherwise Registered + Unregistered
  - rotates the localStorage keys with a fresh version stamp so browsers can
    NEVER show stale auto-saved numbers over a new publish (the fix for the
    "old % still showing in Chrome" problem)

The HTML files stay fully self-contained — safe to email; recipients need
nothing but the one file.
"""

import re
import sys
from datetime import datetime
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
import activity_log  # noqa: E402
import onedrive_paths as op  # noqa: E402

DATA_XLSX = op.DASHBOARDS_DIR / "RCM Dashboard Data.xlsx"
SHEET = "DashboardData"

# metric label in the workbook -> HTML input id (None = handled specially)
METRIC_TO_ID = {
    "Report Date": "in_date",
    "Total Employees": "in_total_emp",
    "Registered": "in_reg",
    "Unregistered": "in_unreg",
    "Training Completed": "in_comp",
    "Trained Denominator (blank = Reg+Unreg)": "in_comp_custom_denom",
    "No Show": "in_noshow",
    # CRB retired 2026-08-21 (column 4 is now Go & Find Status). The HTML
    # keeps these three inputs hidden so these stamps still match; the
    # Go & Find counts are entered in the dashboard itself, not here.
    "CRB Total": "in_crb_tot",
    "CRB Approved": "in_crb_appr",
    "CRB Not Approved": "in_crb_not",
    "All W3 Rev Cycle Total": "in_total_emp_w3",
}

# 2026-08-21: every other HTML dashboard was retired into the Dashboards
# archive folder (archive/retired_2026-08-21) at the analyst's request, so the
# Main (Epic basis) file is the only one still published. The workbook keeps
# its "Excl Lastname09 & Lastname02" column, so bringing that dashboard back is
# just moving the file out of the archive and re-adding its line here.
DASHBOARDS = [
    ("RCM Executive Dashboard W3 Reg.html", 2, "w3reg_epic"),
]


def load_values():
    import openpyxl
    wb = openpyxl.load_workbook(DATA_XLSX, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"ERROR: sheet '{SHEET}' not found in {DATA_XLSX.name}")
    ws = wb[SHEET]
    rows = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] and str(row[0]).strip() in METRIC_TO_ID:
            rows[str(row[0]).strip()] = row
    missing = set(METRIC_TO_ID) - set(rows) - {"CRB Not Approved"}
    if missing:
        sys.exit(f"ERROR: metric row(s) missing from workbook: {sorted(missing)}")
    return rows


def fmt(metric, raw):
    if raw is None or str(raw).strip() == "":
        return None
    if metric == "Report Date":
        if isinstance(raw, datetime):
            return f"{raw:%Y-%m-%d}"
        return str(raw).strip()
    n = float(raw)
    if n != int(n):
        sys.exit(f"ERROR: '{metric}' must be a whole number (got {raw})")
    return str(int(n))


def set_input(html, input_id, val, fname):
    pat = re.compile(r'(id="%s"[^>]*?value=")[^"]*(")' % re.escape(input_id))
    new, n = pat.subn(lambda m: m.group(1) + val + m.group(2), html)
    if n != 1:
        sys.exit(f"ERROR: {fname}: input id={input_id} matched {n} times")
    return new


def set_denom_mode(html, use_custom, fname):
    block_pat = re.compile(r'(<select id="denom_col2".*?</select>)', re.S)
    m = block_pat.search(html)
    if not m:
        sys.exit(f"ERROR: {fname}: denom_col2 select not found")
    block = m.group(1).replace(" selected", "")
    target = 'value="custom_denom"' if use_custom else 'value="reg_pool"'
    block = block.replace(f"<option {target}>", f"<option {target} selected>", 1)
    return html[:m.start(1)] + block + html[m.end(1):]


def force_auto_modes(html, fname):
    """Every donut column must calculate from the counts — a column left in
    'Enter % myself' mode would silently ignore the workbook's numbers
    (this is exactly how a stale 3.3% no-show once shipped)."""
    blocks = re.findall(r'<select id="mode_col\d".*?</select>', html, re.S)
    for b in blocks:
        if 'value="manual" selected' in b:
            nb = (b.replace('<option value="auto">', '<option value="auto" selected>')
                   .replace('value="manual" selected', 'value="manual"'))
            html = html.replace(b, nb)
            print(f"    {fname}: manual override was on — reset to auto")
    return html


def rotate_storage_keys(html, tag, stamp, fname):
    for const in ("SESSION_STORE", "AUTOSAVE_KEY"):
        pat = re.compile(r"(const %s\s*=\s*')[^']*(')" % const)
        prefix = "rcm_sessions_" if const == "SESSION_STORE" else "rcm_autosave_"
        new, n = pat.subn(lambda m: m.group(1) + f"{prefix}{tag}_{stamp}" + m.group(2), html)
        if n != 1:
            sys.exit(f"ERROR: {fname}: {const} matched {n} times")
        html = new
    return html


def main():
    rows = load_values()
    stamp = f"{datetime.now():%Y%m%d%H%M}"
    for fname, col_idx, tag in DASHBOARDS:
        path = op.DASHBOARDS_DIR / fname
        if not path.exists():
            sys.exit(f"ERROR: {path} not found")
        html = path.read_text(encoding="utf-8")
        applied = {}
        for metric, input_id in METRIC_TO_ID.items():
            row = rows.get(metric)
            if row is None:
                continue
            val = fmt(metric, row[col_idx - 1] if len(row) >= col_idx else None)
            if val is None:
                continue
            html = set_input(html, input_id, val, fname)
            applied[metric] = val
        use_custom = "Trained Denominator (blank = Reg+Unreg)" in applied
        html = set_denom_mode(html, use_custom, fname)
        html = force_auto_modes(html, fname)
        html = rotate_storage_keys(html, tag, stamp, fname)
        path.write_text(html, encoding="utf-8")

        reg = int(applied["Registered"]); unreg = int(applied["Unregistered"])
        comp = int(applied["Training Completed"]); ns = int(applied["No Show"])
        pool = reg + unreg
        denom = int(applied["Trained Denominator (blank = Reg+Unreg)"]) if use_custom else pool
        print(f"{fname}")
        print(f"    reg {reg}/{pool} = {reg / pool * 100:.1f}%   "
              f"trained {comp}/{denom} = {comp / denom * 100:.1f}%   "
              f"no-show {ns}/{reg} = {ns / reg * 100:.1f}%   "
              f"(storage keys -> {tag}_{stamp})")


if __name__ == "__main__":
    with activity_log.track_run("refresh_exec_dashboards.py"):
        main()
