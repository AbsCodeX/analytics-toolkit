# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
missing_from_wave_hr.py

READ-ONLY report. HR-side half of the old missing_from_wave.py (now split in
two so the HR-weekly and Epic-daily checks can run on their own schedules —
see missing_from_wave_epic.py for the Epic side).

Lists everyone in the cleaned HR file ("All Staff" tab) who reports to one of
our canonical RCM leaders but is NOT yet in the Master Wave File DATA sheet.

Scope decisions (per ref_reporting_population):
  - "Reports to one of our leaders" = computed Leaders (VP > AVP > SVP
    fallback, canonicalized) lands on a name in the canonical leader list.
    Leaders is RECOMPUTED here from VP/AVP/SVP using the current leader_names
    logic, so it honors the live canonical list (the Leaders column already
    in HR_cleaned is whatever clean_hr.py last wrote and may be stale).
  - Lastname09, Firstname09 and Lastname02, Firstname02 are EXCLUDED. Their orgs are mixed
    RCM + clinical and HR has no Curriculum Type / Job Role to separate them,
    so including them would pull in clinical staff who are out of scope.
    missing_from_wave_epic.py handles those two via its own RCM curriculum
    filter instead.

Output:
    onedrive_paths.MISSING_FROM_WAVE_DIR / hr_missing_from_wave_YYYY.MM.DD.xlsx
    One new dated file per run. Never touches the Master.

Run:
    python scripts/missing_from_wave_hr.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import onedrive_paths
import activity_log
from report_xlsx import write_formatted_xlsx
from leader_names import CANONICAL_SET, NOT_REV_CYCLE, normalize_leader

EXCLUDED_LEADERS = {"Lastname09, Firstname09", "Lastname02, Firstname02"}

OUT_FIELDS = [
    "UniversalID", "Full Name", "Email", "GoLiveWave",
    "Leaders (computed)", "VP", "AVP", "SVP",
    "IsDeparted/Inactive?", "IsRCM",
]


def clean_uid(v):
    if not v:
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.upper() or None


def clean_leader(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.lower() in ("nan", "nan, nan"):
        return None
    return normalize_leader(s)


def compute_leaders(vp, avp, svp):
    """VP > AVP > SVP fallback; non-canonical chosen leader -> No Longer Rev Cycle."""
    raw = vp or avp or svp
    if raw is None:
        return None
    return raw if raw in CANONICAL_SET else NOT_REV_CYCLE


def load_master_uids(path) -> set:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "DATA" not in wb.sheetnames:
        sys.exit("ERROR: 'DATA' sheet not found in the Master Wave File.")
    ws = wb["DATA"]
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    col = {h: i for i, h in enumerate(headers) if h}
    if "UniversalID" not in col:
        sys.exit("ERROR: Master DATA sheet has no 'UniversalID' column.")
    ui = col["UniversalID"]
    uids = set()
    for row in it:
        u = clean_uid(row[ui]) if ui < len(row) else None
        if u:
            uids.add(u)
    wb.close()
    return uids


def build_missing_rows(hr_path, master_uids) -> list[dict]:
    wb = openpyxl.load_workbook(hr_path, read_only=True, data_only=True)
    sheet = "All Staff" if "All Staff" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    col = {h: i for i, h in enumerate(headers) if h}

    def get(row, name):
        i = col.get(name)
        return row[i] if (i is not None and i < len(row)) else None

    seen = set()
    rows_out = []
    for row in it:
        uid = clean_uid(get(row, "UniversalID"))
        if not uid or uid in seen:
            continue
        seen.add(uid)

        vp = clean_leader(get(row, "VP"))
        avp = clean_leader(get(row, "AVP"))
        svp = clean_leader(get(row, "SVP"))
        leaders = compute_leaders(vp, avp, svp)

        if leaders not in CANONICAL_SET:
            continue
        if leaders in EXCLUDED_LEADERS:
            continue
        if uid in master_uids:
            continue

        rows_out.append({
            "UniversalID": uid,
            "Full Name": get(row, "Full Name") or "",
            "Email": get(row, "Email") or "",
            "GoLiveWave": get(row, "GoLiveWave") or "",
            "Leaders (computed)": leaders,
            "VP": vp or "",
            "AVP": avp or "",
            "SVP": svp or "",
            "IsDeparted/Inactive?": get(row, "IsDeparted/Inactive?") or "",
            "IsRCM": get(row, "IsRCM") or "",
        })
    wb.close()

    rows_out.sort(key=lambda r: (r["Leaders (computed)"], str(r["Full Name"])))
    return rows_out


def main() -> int:
    if not onedrive_paths.MASTER_WAVE_PATH.exists():
        print(f"ERROR: Master Wave File not found: {onedrive_paths.MASTER_WAVE_PATH}")
        return 1
    hr_path = onedrive_paths.latest_hr_cleaned()
    if hr_path is None:
        print(f"ERROR: No cleaned HR file found in {onedrive_paths.CLEAN_HR_DIR}  (run clean_hr.py first)")
        return 1

    master_uids = load_master_uids(onedrive_paths.MASTER_WAVE_PATH)
    print(f"Master DATA UIDs: {len(master_uids):,}")
    print(f"Loading HR file: {hr_path.name}")
    rows = build_missing_rows(hr_path, master_uids)

    out_path = onedrive_paths.MISSING_FROM_WAVE_DIR / f"{onedrive_paths.dated_stem('hr_missing_from_wave')}.xlsx"
    write_formatted_xlsx(pd.DataFrame(rows, columns=OUT_FIELDS), out_path, "Missing From HR")

    from collections import Counter
    by_leader = Counter(r["Leaders (computed)"] for r in rows)
    departed = sum(1 for r in rows
                   if str(r["IsDeparted/Inactive?"]).strip().lower() in ("true", "yes", "1", "y"))
    print(f"\nMissing From HR (in scope): {len(rows):,}")
    print(f"  of which flagged departed/inactive: {departed:,}")
    print("  (Lastname09, Firstname09 + Lastname02, Firstname02 excluded — mixed RCM/clinical)")
    print("By leader:")
    for ldr, n in by_leader.most_common():
        print(f"  {n:>4}  {ldr}")
    print(f"\nWritten: {out_path}")
    return 0


if __name__ == "__main__":
    with activity_log.track_run("missing_from_wave_hr.py"):
        sys.exit(main())
