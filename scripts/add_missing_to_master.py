# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
add_missing_to_master.py — add missing people to the Master Wave File from
TWO sources (OneDrive-era successor to add_new_centralized_members.py; same
safe xlwings Table-append pattern, new inputs):

  EPIC  : Epic Team Member Lookup.xlsx (Export sheet), Curriculum Type
          "Revenue Cycle - Centralized", Universal ID not in Master.
  HR    : HR_cleaned All Staff, SVP = Lastname03, Firstname03, not departed/inactive,
          UniversalID not in Master.

Enrichment precedence: leader chain / Leaders / GoLiveWave from HR_cleaned
All Staff (already resolved + canonicalized by clean_hr.py); identity detail
(names, EmployeeID, JobTitle, Department, manager UIDs) from raw HR.xlsx,
falling back to the Epic export for Epic-sourced people.

SAFE-WRITE PATTERN (unchanged from add_new_centralized_members.py):
  * Master is READ with openpyxl but only ever WRITTEN through desktop Excel
    (xlwings/COM) via the DATA sheet's ListObject ("Table1") ListRows.Add(),
    so the table extends and calculated columns auto-fill.
  * Formula columns are never written. Existing rows are never touched.
  * Daily full backup + values-only DATA snapshot via master_backup.
  * Refuses to run if the Master is open/locked in Excel.
  * DRY RUN by default; review CSV is written in both modes.

USAGE
-----
    python scripts/add_missing_to_master.py            # DRY RUN preview
    python scripts/add_missing_to_master.py --apply    # backup + append + log
"""

import argparse
import datetime
import sys
import warnings
from pathlib import Path

import openpyxl  # READS only — never used to write the Master
import pandas as pd

warnings.filterwarnings("ignore")

sys.path.insert(0, str(Path(__file__).parent))
import master_backup
import onedrive_paths as op
import runlog
from activity_log import track_run
from leader_names import CANONICAL_SET, NOT_REV_CYCLE, normalize_leader

TODAY = datetime.date.today()
TODAY_SHORT = TODAY.strftime("%m/%d/%y")

MASTER_PATH = op.MASTER_WAVE_PATH
EPIC_PATH = op.RAW_EPIC_TM_DIR / "Epic Team Member Lookup.xlsx"
HR_RAW_PATH = op.RAW_HR_PATH
# dated review CSVs live in their own folder since the 2026-07-27 reports
# cleanup (they used to pile up loose next to the Master)
CSV_OUT = MASTER_PATH.parent / "new_members_added" / f"New_Members_Added_{TODAY.isoformat()}.csv"

TARGET_CURRICULUM = "Revenue Cycle - Centralized"
TARGET_SVP = "Lastname03, Firstname03"
DATA_SHEET = "DATA"
TABLE_NAME = "Table1"
FORMULA_COLS = {"Vendor Yes/No", "Leader? (Yes/No)"}
# If the combined add-list exceeds this, something upstream is wrong (an empty
# Master read, a malformed export) — refuse to write and ask for eyes.
SANITY_MAX_ADDS = 150

RUNLOG_TAB = "New Members Added"
RUNLOG_HEADERS = [
    "Run Timestamp", "Source Exports", "Candidates Scanned",
    "Already in Master", "New Members Found", "New Members Added",
    "Wave Breakdown", "UIDs Added", "Master File", "Backup File",
]


def norm_uid(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.upper() or None


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.lower() in ("nan", "nan, nan", "none"):
        return None
    return s


def parse_last_first(name):
    s = clean(name)
    if not s:
        return None, None
    if "," in s:
        last, first = s.split(",", 1)
        return clean(first), clean(last)
    return None, s


def norm_wave(v):
    s = clean(v)
    if not s:
        return None
    if s.lower().startswith("wave"):
        return s
    if s.isdigit():
        return f"Wave {s}"
    return s


def compute_leaders(vp, avp, svp):
    raw = vp or avp or svp
    if raw is None:
        return None
    return raw if raw in CANONICAL_SET else NOT_REV_CYCLE


# ---------------------------------------------------------------------------
# Loads (all READ ONLY)
# ---------------------------------------------------------------------------
def load_master_uids():
    if office_lock_present(MASTER_PATH):
        sys.exit("ERROR: Master Wave File is open/locked in Excel. Close it and re-run.")
    try:
        wb = openpyxl.load_workbook(MASTER_PATH, read_only=True, data_only=True)
    except PermissionError:
        sys.exit("ERROR: Master Wave File is locked (open in Excel or mid-OneDrive-sync). "
                 "Close it and re-run.")
    ws = wb[DATA_SHEET]
    headers = [str(c.value).strip() if c.value is not None else ""
               for c in next(ws.iter_rows(max_row=1))]
    uid_i = headers.index("UniversalID")
    uids = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        u = norm_uid(row[uid_i]) if uid_i < len(row) else None
        if u:
            uids.add(u)
    wb.close()
    if not uids:
        sys.exit("ERROR: 0 UIDs read from Master — refusing to continue.")
    return uids


def load_epic_centralized():
    if not EPIC_PATH.exists():
        sys.exit(f"ERROR: Epic export not found: {EPIC_PATH}")
    df = pd.read_excel(EPIC_PATH, sheet_name="Export", engine="openpyxl", dtype=str)
    for col in ("Universal ID", "Curriculum Type", "Team Member"):
        if col not in df.columns:
            sys.exit(f"ERROR: Epic Export sheet is missing column '{col}'.")
    df = df[df["Curriculum Type"].astype(str).str.strip() == TARGET_CURRICULUM].copy()
    df["UID"] = df["Universal ID"].map(norm_uid)
    df = df.dropna(subset=["UID"]).drop_duplicates(subset=["UID"])
    if df.empty:
        sys.exit("ERROR: 0 Centralized rows in the Epic export — refusing to continue.")
    return df.set_index("UID")


def load_hr_cleaned_all_staff():
    path = op.latest_hr_cleaned()
    if path is None:
        sys.exit("ERROR: no HR_cleaned_*.xlsx found — run clean_hr.py first.")
    df = pd.read_excel(path, sheet_name="All Staff", engine="openpyxl", dtype=str)
    for col in ("UniversalID", "SVP", "Leaders", "Full Name"):
        if col not in df.columns:
            sys.exit(f"ERROR: '{path.name}' All Staff is missing column '{col}'.")
    df["UID"] = df["UniversalID"].map(norm_uid)
    df = df.dropna(subset=["UID"]).drop_duplicates(subset=["UID"])
    return df.set_index("UID"), path


def load_hr_raw():
    """Raw HR keyed by UID for identity detail (names, EmployeeID, JobTitle...)."""
    if not HR_RAW_PATH.exists():
        return pd.DataFrame().set_index(pd.Index([], name="UID")), None
    sheet = openpyxl.load_workbook(HR_RAW_PATH, read_only=True).sheetnames[0]
    hr = pd.read_excel(HR_RAW_PATH, sheet_name=sheet, engine="openpyxl", dtype=str)
    hr["UID"] = hr["UniversalID"].map(norm_uid)
    hr = hr.dropna(subset=["UID"]).drop_duplicates(subset=["UID"])
    return hr.set_index("UID"), sheet


def is_true(v):
    return clean(v) is not None and str(v).strip().lower() in ("true", "yes", "1", "y")


# ---------------------------------------------------------------------------
# Row builder
# ---------------------------------------------------------------------------
def build_row(uid, source, epic, staff, hr_raw):
    e = epic.loc[uid].to_dict() if uid in epic.index else {}
    s = staff.loc[uid].to_dict() if uid in staff.index else {}
    h = hr_raw.loc[uid].to_dict() if uid in hr_raw.index else {}

    e_first, e_last = parse_last_first(e.get("Team Member"))
    first = clean(h.get("FirstName")) or e_first
    last = clean(h.get("LastName")) or e_last
    mid = clean(h.get("MiddleName"))
    mi = (mid[0] + ".") if mid else None
    full = (clean(s.get("Full Name")) or clean(e.get("Team Member"))
            or (f"{last}, {first}" if last and first else (last or first)))

    # Chain: HR_cleaned All Staff is authoritative (already name-resolved).
    # Leaders is recomputed against the canonical allowlist (VP > AVP > SVP,
    # else No Longer Rev Cycle) — same rule as the Master writers — so
    # out-of-scope chains land in review instead of as stray leader names.
    vp = normalize_leader(clean(s.get("VP")))
    avp = normalize_leader(clean(s.get("AVP")))
    svp = normalize_leader(clean(s.get("SVP")))
    leaders = compute_leaders(vp, avp, svp)

    wave = norm_wave(s.get("GoLiveWave")) or norm_wave(e.get("Wave"))

    in_hr = uid in staff.index
    note = (f"Added {TODAY_SHORT} - missing from wave file ({source}); "
            f"not previously in Master.")
    if not in_hr:
        note += " NOT FOUND IN HR — verify manually."
    wave_log = (f"{TODAY_SHORT} - Added to Master ({source} missing-from-wave); "
                f"wave + leadership seeded from HR_cleaned/Epic.")

    return {
        "UniversalID": uid,
        "EmployeeID": clean(h.get("EmployeeID")) or clean(e.get("Employee ID")),
        "FirstName": first,
        "MI": mi,
        "LastName": last,
        "Full Name": full,
        "DepartmentLocation": clean(h.get("Department Name")) or clean(e.get("Department")),
        "JobTitle": clean(h.get("JobTitle")) or clean(e.get("Job Title")),
        "ValidatingManagerUniversalID": norm_uid(h.get("ManagerUniversalID"))
            or norm_uid(e.get("Validating Manager Universal ID")),
        "ValidatingManagerFirstName": clean(h.get("ManagerFirstName")),
        "ValidatingManagerLastName": clean(h.get("ManagerLastName")),
        "HRManagerUniversalID": norm_uid(h.get("ManagerUniversalID")),
        "Status": "Not Started",
        "GoLiveWave": wave,
        "IsDepartedInactive?": False,
        "Badge Buddies": False,
        "Training Needed (Yes/No)": "Yes",
        "SeniorManager": clean(s.get("SeniorManager")),
        "Director": clean(s.get("Director")),
        "Sr. Director": clean(s.get("SeniorDirector")),
        "AVP": avp,
        "VP": vp,
        "Leaders": leaders,
        "SVP": svp,
        "Notes": note,
        "Users Wave Change Log": wave_log,
    }


def office_lock_present(path):
    return path.with_name("~$" + path.name).exists()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main(apply):
    banner = "LIVE RUN — WILL APPEND ROWS" if apply else "DRY RUN — NOTHING WILL BE WRITTEN"
    print("=" * 72)
    print(f"  add_missing_to_master.py   [{banner}]")
    print("=" * 72 + "\n")

    master_uids = load_master_uids()
    print(f"Master UIDs: {len(master_uids):,}")

    epic = load_epic_centralized()
    epic_new = [u for u in epic.index if u not in master_uids]
    print(f"Epic '{TARGET_CURRICULUM}': {len(epic):,} people | "
          f"already in Master: {len(epic) - len(epic_new):,} | NEW: {len(epic_new)}")

    staff, staff_path = load_hr_cleaned_all_staff()
    lastname03 = staff[staff["SVP"].astype(str).str.strip() == TARGET_SVP]
    brogan_active = lastname03[~lastname03.get("IsDeparted/Inactive?", "").map(is_true)]
    hr_new = [u for u in brogan_active.index if u not in master_uids and u not in set(epic_new)]
    print(f"HR All Staff ({staff_path.name}) SVP={TARGET_SVP}: {len(lastname03):,} people "
          f"({len(lastname03) - len(brogan_active)} departed excluded) | NEW: {len(hr_new)}\n")

    additions = [(u, "Epic Centralized") for u in sorted(epic_new)] + \
                [(u, "HR Lastname03 org") for u in sorted(hr_new)]
    if not additions:
        print("Nothing to add — no one is missing from the Master.")
        return

    if len(additions) > SANITY_MAX_ADDS:
        sys.exit(f"SANITY ABORT: {len(additions)} candidates exceeds {SANITY_MAX_ADDS} — "
                 "an upstream file looks wrong. No changes made.")

    hr_raw, _hr_sheet = load_hr_raw()
    rows = [build_row(u, src, epic, staff, hr_raw) for u, src in additions]
    out_df = pd.DataFrame(rows)
    out_df.insert(1, "Source", [src for _, src in additions])
    out_df.to_csv(CSV_OUT, index=False)

    show = ["UniversalID", "Source", "Full Name", "GoLiveWave", "JobTitle", "Leaders"]
    pd.set_option("display.max_rows", None, "display.width", 220, "display.max_colwidth", 30)
    print("Rows to add:")
    print(out_df[show].to_string(index=False))
    wave_counts = out_df["GoLiveWave"].value_counts(dropna=False).to_dict()
    print(f"\n  Wave breakdown: {wave_counts}")
    print(f"  Review CSV: {CSV_OUT}\n")

    if not apply:
        print("DRY RUN — no backup made, nothing written. Re-run with --apply.")
        return

    import xlwings as xw

    if office_lock_present(MASTER_PATH):
        sys.exit("ERROR: Master Wave File is open/locked in Excel. Close it and re-run.")

    backup_path = master_backup.ensure_daily_full_backup(MASTER_PATH)
    print(f"Backup: {backup_path}")

    drop_source = out_df.drop(columns=["Source"])
    row_dicts = drop_source.where(pd.notna(drop_source), None).to_dict("records")

    app = wb = None
    orig_calc = orig_alerts = orig_screen = None
    created = False
    try:
        app = xw.App(visible=False, add_book=False)
        created = True
        orig_alerts, orig_screen = app.display_alerts, app.screen_updating
        app.display_alerts = app.screen_updating = False

        wb = app.books.open(str(MASTER_PATH), update_links=False)
        orig_calc = app.calculation
        app.calculation = "manual"

        ws = wb.sheets[DATA_SHEET]
        lo = ws.tables[TABLE_NAME].api

        header_cells = lo.HeaderRowRange.Value[0]
        col_idx = {str(h).strip(): i + 1 for i, h in enumerate(header_cells) if h is not None}
        missing = sorted({c for r in row_dicts for c in r if c not in col_idx})
        if missing:
            sys.exit(f"ERROR: target column(s) not found in DATA table: {missing}")

        first_new = lo.Range.Row + lo.Range.Rows.Count
        print(f"Appending {len(row_dicts)} row(s) starting at sheet row {first_new}...")
        for i, rowvals in enumerate(row_dicts):
            lo.ListRows.Add()
            r = first_new + i
            for header, val in rowvals.items():
                if header in FORMULA_COLS or val is None:
                    continue
                ws.range((r, col_idx[header])).value = val

        app.calculation = orig_calc
        orig_calc = None
        wb.save()
        print("Saved through Excel (workbook package + Table preserved).")
    finally:
        try:
            if app is not None and orig_calc is not None:
                app.calculation = orig_calc
        except Exception:
            pass
        try:
            if wb is not None:
                wb.close()
        except Exception:
            pass
        try:
            if app is not None:
                app.display_alerts = True if orig_alerts is None else orig_alerts
                app.screen_updating = True if orig_screen is None else orig_screen
        except Exception:
            pass
        try:
            if created and app is not None:
                app.quit()
        except Exception:
            pass

    snap = master_backup.save_data_snapshot(MASTER_PATH)
    print(f"DATA snapshot: {snap}")

    wave_str = "; ".join(f"{k}={v}" for k, v in wave_counts.items())
    runlog.append(RUNLOG_TAB, RUNLOG_HEADERS, [
        datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Epic TM Lookup (Centralized) + HR_cleaned (SVP Lastname03)",
        len(epic) + len(brogan_active),
        (len(epic) - len(epic_new)) + (len(brogan_active) - len(hr_new)),
        len(additions), len(additions), wave_str,
        ", ".join(out_df["UniversalID"].tolist()),
        str(MASTER_PATH), str(backup_path),
    ])
    print(f"\nAdded {len(additions)} new member(s) "
          f"({len(epic_new)} Epic Centralized, {len(hr_new)} HR Lastname03 org).")
    print("Next: run the daily apply so MVP/Epic updates + reports pick them up.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true",
                    help="Back up the Master and append. Default is a dry-run preview.")
    with track_run("Add Missing to Master"):
        main(apply=ap.parse_args().apply)
