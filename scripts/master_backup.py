# =========================================================================
# SANITIZED TEMPLATE — org-specific values were replaced with placeholders.
# Search for: YourOrg, YourUser, Lastname01..NN, C:\path\to\...  and
# fill in your own values before running. See README.md for the full map.
# =========================================================================
"""
master_backup.py

Backup/versioning helpers for the Master Wave File, shared by
update_master_from_mvp.py and update_master_from_hr_safe.py.

Two mechanisms, deliberately different cadences:

1. ensure_daily_full_backup() — a full-workbook copy, AT MOST ONCE PER
   CALENDAR DAY (skipped if today's backup already exists). This replaced an
   earlier pattern of copying the whole ~5-7MB workbook on every single run,
   which produced 3 backups (~18MB) in one afternoon of testing. One full
   backup per day is still a genuine pre-change safety net (taken before that
   day's first edit) without the waste.

2. save_data_snapshot() — a lightweight, VALUES-ONLY copy of just the DATA
   sheet (no formulas, no other sheets — ~10x smaller than a full-workbook
   copy), taken after EVERY successful Master edit. This is the versioned,
   git-like history of what DATA actually looked like after each run.

Both land in monthly subfolders (onedrive_paths.month_subdir) so the archive
folders stay browsable as they grow — retention is intentionally unbounded
(the user asked not to prune), just organized.
"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

import onedrive_paths


def ensure_daily_full_backup(master_path: Path) -> Path:
    """Copy the full Master workbook to MASTER_ARCHIVE_DIR/<month>/ if today's
    backup doesn't already exist (no-op otherwise). Always returns today's
    backup path — whether just-created or already there — so callers always
    have somewhere to point a "restore from backup" message."""
    month_dir = onedrive_paths.month_subdir(onedrive_paths.MASTER_ARCHIVE_DIR)
    today_stamp = datetime.now().strftime("%Y-%m-%d")
    backup_path = month_dir / f"{master_path.stem}.bak.{today_stamp}.xlsx"
    if backup_path.exists():
        print(f"  Today's full backup already exists — skipped: {backup_path.name}")
    else:
        shutil.copy(master_path, backup_path)
        print(f"  Full backup created: {backup_path.name}")
    return backup_path


def save_data_snapshot(master_path: Path, sheet_name: str = "DATA") -> Path | None:
    """Write a values-only copy of the Master's DATA sheet to
    WAVE_REPOSITORY_DIR/<month>/DATA_<timestamp>.xlsx. Unlike the daily full
    backup, this runs once per actual edit (multiple per day is expected).

    Best-effort: this is called AFTER the Master has already been saved, so a
    snapshot failure must never make an otherwise-successful run look failed.
    On any error it warns and returns None instead of raising.
    """
    try:
        src = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
        try:
            ws = src[sheet_name]
            month_dir = onedrive_paths.month_subdir(onedrive_paths.WAVE_REPOSITORY_DIR)
            stamp = datetime.now().strftime("%Y.%m.%d_%H%M%S")
            dest_path = month_dir / f"DATA_{stamp}.xlsx"

            out = openpyxl.Workbook(write_only=True)
            out_ws = out.create_sheet(sheet_name)
            for row in ws.iter_rows(values_only=True):
                out_ws.append(row)
            out.save(dest_path)
        finally:
            src.close()
        return dest_path
    except Exception as e:
        print(f"  WARNING: DATA snapshot could not be written ({type(e).__name__}: {e}). "
              "The Master save itself was not affected.")
        return None
